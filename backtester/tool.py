
import os
from docx import Document
import openpyxl

def count_words_in_docx(docx_path):
    document = Document(docx_path)
    word_count = 0
    for paragraph in document.paragraphs:
        word_count += len(paragraph.text.split())
    return word_count

def process_folder(folder_path, excel_file_path):
    # 创建Excel文件和工作表
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Word Count"

    # 添加表头
    ws.append(["文件名", "字数", "出处"])

    # 遍历文件夹中的所有文件
    for root, dirs, files in os.walk(folder_path):
        for file_name in files:
            if file_name.endswith(".docx"):
                file_path = os.path.join(root, file_name)
                word_count = count_words_in_docx(file_path)

                # 在Excel表格中添加数据
                ws.append([file_name, word_count, file_path])

    # 保存Excel文件
    wb.save(excel_file_path)
    print(f"数据已保存至 {excel_file_path}")

if __name__ == "__main__":
    folder_path = "C:/Users/Administrator/Desktop/test"  # 替换为你的实际文件夹路径
    excel_file_path = "test.xlsx"  # 替换为你想要保存的Excel文件路径

    process_folder(folder_path, excel_file_path)
